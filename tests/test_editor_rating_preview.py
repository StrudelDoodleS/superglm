"""The rating-table preview shows the Excel export's own block for the current term."""

from __future__ import annotations

import io
import json
import urllib.request

import numpy as np
import pandas as pd
import pytest

from superglm import (
    Categorical,
    LambdaPolicy,
    Numeric,
    OrderedCategorical,
    Piecewise,
    RandomEffect,
    Spline,
    SuperGLM,
)
from superglm.editor import EditorSession
from superglm.editor.rating_preview import EXPORT_REFUSED, UNSUPPORTED_TERMS

TERMS = ["x_spline", "x_piece", "region", "band"]


def _frame(seed=20261003, n=600):
    rng = np.random.default_rng(seed)
    X = pd.DataFrame(
        {
            "x_spline": rng.uniform(0.0, 10.0, n),
            "x_piece": rng.uniform(0.0, 6.0, n),
            "region": rng.choice(["A", "B", "C"], n),
            "band": rng.choice(["low", "medium", "high"], n),
        }
    )
    eta = (
        -0.6
        + 0.18 * np.sin(X["x_spline"].to_numpy() / 1.8)
        + 0.05 * np.abs(X["x_piece"].to_numpy() - 3.0)
        + 0.25 * (X["region"].to_numpy() == "B")
        + 0.3 * (X["band"].to_numpy() == "high")
    )
    return X, rng.poisson(np.exp(eta)).astype(np.float64)


def _features():
    return {
        "x_spline": Spline(n_knots=6),
        "x_piece": Piecewise(breaks=[3.0]),
        "region": Categorical(base="first"),
        "band": OrderedCategorical(
            order=["low", "medium", "high"], basis=Spline(kind="ps", n_knots=2), base="first"
        ),
    }


@pytest.fixture
def poisson_session():
    X, y = _frame()
    model = SuperGLM(family="poisson", selection_penalty=0.0, features=_features()).fit(X, y)
    return EditorSession.from_model(model, terms=TERMS, train_data=(X, y))


def _workbook_block(data: bytes, term: str):
    """``(columns, rows, formats, note)`` of ``term``'s block on the workbook's sheet."""
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data))["Rating Tables"]
    start = next(cell.column for cell in sheet[5] if cell.value == term)
    width = 0
    while sheet.cell(row=7, column=start + width).value not in (None, ""):
        width += 1
        if sheet.cell(row=5, column=start + width).value not in (None, ""):
            break
    columns = [sheet.cell(row=7, column=start + offset).value for offset in range(width)]
    rows, formats, row = [], None, 8
    while sheet.cell(row=row, column=start).value is not None:
        cells = [sheet.cell(row=row, column=start + offset) for offset in range(width)]
        rows.append([cell.value for cell in cells])
        formats = [
            None if cell.number_format == "General" else cell.number_format for cell in cells
        ]
        row += 1
    return columns, rows, formats, sheet.cell(row=6, column=start).value


def _assert_same_cells(sent, written):
    """Text cells match exactly; numbers to the workbook's 16 significant digits.

    openpyxl writes a float as ``"%.16g"``, half a unit in the 16th digit at
    most ``5e-16 |v|`` away, and reading the decimal back rounds once more, by
    ``u`` relative. The preview sends the block's float64 values themselves.
    """
    assert len(sent) == len(written)
    for value, cell in zip(sent, written, strict=True):
        if isinstance(value, str):
            assert value == cell
        else:
            assert abs(value - cell) <= 5e-16 * abs(value) + 2.0**-53 * abs(cell)


def test_the_preview_is_the_workbook_block_for_every_term(poisson_session):
    # On master there is no preview at all: the widget has no `_rating_table`.
    session = poisson_session
    session.select_levels("region", ["B"])
    session.shift("region", 0.2)
    widget = session.widget()
    try:
        workbook = widget._export_bytes("xlsx").data
        previews = {term: widget._rating_table(term) for term in TERMS}
    finally:
        widget.close()

    for term, preview in previews.items():
        columns, rows, formats, note = _workbook_block(workbook, term)
        assert preview["available"] is True
        assert preview["model_revision"] == session.model_revision
        assert preview["columns"] == columns
        assert len(preview["rows"]) == len(rows)
        for sent, written in zip(preview["rows"], rows, strict=True):
            _assert_same_cells(sent, written)
        assert preview["formats"] == formats
        assert preview["note"] == note


def test_the_preview_builds_once_per_revision_and_skips_the_impact_sweep(
    poisson_session, monkeypatch
):
    from superglm.export import rating_tables

    calls = []
    build = rating_tables.build_rating_table_payload

    def counting(*args, **kwargs):
        calls.append(kwargs.get("impact_bins"))
        return build(*args, **kwargs)

    monkeypatch.setattr(rating_tables, "build_rating_table_payload", counting)
    session = poisson_session
    widget = session.widget()
    try:
        widget._rating_table("region")
        widget._rating_table("band")
        assert calls == [()]
        session.select_levels("band", ["high"])
        session.shift("band", 0.1)
        edited = widget._rating_table("band")
        assert calls == [(), ()]
    finally:
        widget.close()
    assert edited["model_revision"] == session.model_revision


def test_without_training_data_the_preview_gives_the_exports_sentence():
    rng = np.random.default_rng(20260802)
    X = pd.DataFrame({"x": rng.normal(size=90)})
    y = rng.poisson(np.exp(0.2 + 0.4 * X["x"].to_numpy())).astype(np.float64)
    model = SuperGLM(
        family="poisson",
        retain_fit_state=False,
        selection_penalty=0.0,
        features={"x": Numeric()},
    ).fit(X, y)
    session = EditorSession.from_model(model, terms=["x"], validation_data=(X[:20], y[:20]))
    widget = session.widget()
    try:
        preview = widget._rating_table("x")
    finally:
        widget.close()

    assert preview["available"] is False
    assert preview["reason"] == (
        "Excel export requires train_data or retained fit data; "
        "validation/test data are not substituted."
    )
    assert preview["rows"] == []


def test_a_refused_export_gives_a_fixed_sentence_not_backend_text():
    # A gaussian identity-link model: the multiplicative workbook refuses it.
    X, y = _frame()
    model = SuperGLM(family="gaussian", selection_penalty=0.0, features=_features()).fit(
        X, np.log1p(y)
    )
    session = EditorSession.from_model(model, terms=TERMS)
    widget = session.widget()
    try:
        preview = widget._rating_table("region")
    finally:
        widget.close()

    assert preview["available"] is False
    assert preview["reason"] == EXPORT_REFUSED


def test_a_model_with_a_random_effect_says_rating_tables_do_not_cover_it():
    rng = np.random.default_rng(20260726)
    codes = np.repeat(np.arange(8), 12)
    X = pd.DataFrame({"x": rng.normal(size=codes.size), "group": [f"g{c}" for c in codes]})
    y = rng.poisson(np.exp(0.1 * X["x"].to_numpy() + 0.05 * codes)).astype(np.float64)
    model = SuperGLM(
        family="poisson",
        features={"x": Numeric(), "group": RandomEffect(lambda_policy=LambdaPolicy.fixed(1.0))},
        selection_penalty=0.0,
        direct_solve="structured",
    ).fit_reml(X, y, runtime_validation="skip")
    session = EditorSession.from_model(model, terms=["x"])
    widget = session.widget()
    try:
        preview = widget._rating_table("x")
    finally:
        widget.close()

    assert preview["available"] is False
    assert preview["reason"] == UNSUPPORTED_TERMS


def test_the_rating_table_route_answers_in_the_contract_shape(poisson_session):
    widget = poisson_session.widget()
    try:
        request = urllib.request.Request(
            f"{widget.url}/rating_table",
            data=json.dumps({"term": "band"}).encode("utf-8"),
            method="POST",
            headers={
                "Content-Type": "application/json",
                "X-SuperGLM-Editor-Token": widget._token,
            },
        )
        with urllib.request.urlopen(request, timeout=30) as response:
            payload = json.loads(response.read().decode("utf-8"))
    finally:
        widget.close()

    assert set(payload) == {
        "term",
        "available",
        "reason",
        "columns",
        "rows",
        "formats",
        "note",
        "model_revision",
    }
    assert payload["term"] == "band"
    assert payload["columns"] == ["band", "Relativity", "Weight"]
    assert [row[0] for row in payload["rows"]] == ["low", "medium", "high"]
